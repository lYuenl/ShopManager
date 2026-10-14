from __future__ import annotations

import asyncio
import os
import socket
import sqlite3
from collections import defaultdict
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from nicegui import app as web_app, ui
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / 'data'
EXPORT_DIR = BASE_DIR / 'exports'
BACKUP_DIR = BASE_DIR / 'backups'
DB_PATH = DATA_DIR / 'shop.sqlite3'
PORT = int(os.getenv('SHOP_PORT', '8080'))

for directory in (DATA_DIR, EXPORT_DIR, BACKUP_DIR):
    directory.mkdir(parents=True, exist_ok=True)

DEFAULT_SHOP_NAME = '我的店家'
DEFAULT_SHOP_SUBTITLE = '營運管理系統'

INCOME_CATEGORIES = ['商品銷售', '其他收入']
EXPENSE_CATEGORIES = ['原料進貨', '包裝耗材', '水電瓦斯', '租金', '設備維修', '運費', '人事費用', '退款／退費', '其他支出']
PAYMENT_METHODS = ['現金', 'LINE Pay', '信用卡／電子支付', '轉帳', '其他']


def db_connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys = ON')
    return conn


@contextmanager
def db_session():
    conn = db_connect()
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def init_db() -> None:
    with db_session() as conn:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS products (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                sale_price INTEGER NOT NULL DEFAULT 0,
                unit_cost INTEGER NOT NULL DEFAULT 0,
                active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                transacted_at TEXT NOT NULL,
                type TEXT NOT NULL CHECK(type IN ('income', 'expense')),
                category TEXT NOT NULL,
                description TEXT NOT NULL DEFAULT '',
                amount INTEGER NOT NULL CHECK(amount >= 0),
                payment_method TEXT NOT NULL DEFAULT '現金',
                product_id INTEGER,
                product_name_snapshot TEXT NOT NULL DEFAULT '',
                quantity REAL NOT NULL DEFAULT 0,
                unit_cost_snapshot INTEGER NOT NULL DEFAULT 0,
                notes TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                FOREIGN KEY(product_id) REFERENCES products(id) ON DELETE SET NULL
            )
        ''')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_transactions_date ON transactions(transacted_at)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_transactions_type ON transactions(type)')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL DEFAULT ''
            )
        ''')


init_db()


def get_setting(key: str, default: str = '') -> str:
    with db_session() as conn:
        row = conn.execute('SELECT value FROM settings WHERE key=?', (key,)).fetchone()
    return row['value'] if row else default


def set_setting(key: str, value: str) -> None:
    with db_session() as conn:
        conn.execute(
            'INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value',
            (key, value),
        )


def get_shop_name() -> str:
    return get_setting('shop_name', DEFAULT_SHOP_NAME) or DEFAULT_SHOP_NAME


def get_shop_subtitle() -> str:
    return get_setting('shop_subtitle', DEFAULT_SHOP_SUBTITLE) or DEFAULT_SHOP_SUBTITLE


def money(value: int | float | None) -> str:
    return f"${int(value or 0):,}"


def parse_dt(value: str) -> datetime:
    return datetime.fromisoformat(value)


def get_products(active_only: bool = False) -> list[dict[str, Any]]:
    sql = 'SELECT * FROM products' + (' WHERE active=1' if active_only else '') + ' ORDER BY active DESC, name'
    with db_session() as conn:
        return [dict(row) for row in conn.execute(sql).fetchall()]


def get_transaction(txn_id: int) -> dict[str, Any] | None:
    with db_session() as conn:
        row = conn.execute('SELECT * FROM transactions WHERE id=?', (txn_id,)).fetchone()
        return dict(row) if row else None


def get_transactions(start: datetime | None = None, end: datetime | None = None, limit: int | None = None) -> list[dict[str, Any]]:
    clauses: list[str] = []
    params: list[Any] = []
    if start is not None:
        clauses.append('transacted_at >= ?')
        params.append(start.isoformat(timespec='minutes'))
    if end is not None:
        clauses.append('transacted_at <= ?')
        params.append(end.isoformat(timespec='minutes'))
    sql = 'SELECT * FROM transactions'
    if clauses:
        sql += ' WHERE ' + ' AND '.join(clauses)
    sql += ' ORDER BY transacted_at DESC, id DESC'
    if limit is not None:
        sql += ' LIMIT ?'
        params.append(limit)
    with db_session() as conn:
        return [dict(row) for row in conn.execute(sql, params).fetchall()]


def sums_for_range(start: datetime, end: datetime) -> dict[str, int]:
    with db_session() as conn:
        rows = conn.execute(
            '''SELECT type, COALESCE(SUM(amount),0) AS total
               FROM transactions WHERE transacted_at >= ? AND transacted_at <= ? GROUP BY type''',
            (start.isoformat(timespec='minutes'), end.isoformat(timespec='minutes')),
        ).fetchall()
    result = {'income': 0, 'expense': 0}
    for row in rows:
        result[row['type']] = int(row['total'] or 0)
    return result


def period_bounds(period: str, start_text: str | None = None, end_text: str | None = None) -> tuple[datetime, datetime, str]:
    now = datetime.now().replace(second=0, microsecond=0)
    if period == '24h':
        return now - timedelta(hours=24), now, '最近 24 小時'
    if period == '7d':
        start = (now - timedelta(days=6)).replace(hour=0, minute=0)
        return start, now, '最近 7 天'
    if period == '30d':
        start = (now - timedelta(days=29)).replace(hour=0, minute=0)
        return start, now, '最近 30 天'
    if period == 'month':
        return now.replace(day=1, hour=0, minute=0), now, '本月'
    if period == 'year':
        return now.replace(month=1, day=1, hour=0, minute=0), now, '今年'
    if period == 'custom':
        try:
            start_day = date.fromisoformat(start_text or date.today().isoformat())
            end_day = date.fromisoformat(end_text or date.today().isoformat())
        except ValueError:
            start_day = end_day = date.today()
        if end_day < start_day:
            start_day, end_day = end_day, start_day
        start_dt = datetime.combine(start_day, datetime.min.time())
        end_dt = datetime.combine(end_day, datetime.max.time()).replace(second=0, microsecond=0)
        return start_dt, end_dt, f'{start_day.isoformat()} ～ {end_day.isoformat()}'
    return (now - timedelta(days=6)).replace(hour=0, minute=0), now, '最近 7 天'


def chart_series(period: str, start: datetime, end: datetime) -> tuple[list[str], list[int], list[int], list[int]]:
    txns = get_transactions(start, end)
    revenue: defaultdict[str, int] = defaultdict(int)
    expenses: defaultdict[str, int] = defaultdict(int)
    if period == '24h':
        first_hour = start.replace(minute=0)
        labels = [(first_hour + timedelta(hours=i)).strftime('%H:%M') for i in range(25)]
        bucket_keys = [(first_hour + timedelta(hours=i)).strftime('%Y-%m-%dT%H') for i in range(25)]
        for txn in txns:
            key = parse_dt(txn['transacted_at']).strftime('%Y-%m-%dT%H')
            (revenue if txn['type'] == 'income' else expenses)[key] += txn['amount']
    elif period == 'year':
        labels = [f'{i}月' for i in range(1, 13)]
        bucket_keys = [f'{start.year}-{i:02d}' for i in range(1, 13)]
        for txn in txns:
            dt = parse_dt(txn['transacted_at'])
            key = dt.strftime('%Y-%m')
            (revenue if txn['type'] == 'income' else expenses)[key] += txn['amount']
    else:
        days = (end.date() - start.date()).days
        bucket_dates = [start.date() + timedelta(days=i) for i in range(days + 1)]
        bucket_keys = [d.isoformat() for d in bucket_dates]
        labels = [d.strftime('%m/%d') for d in bucket_dates]
        for txn in txns:
            key = parse_dt(txn['transacted_at']).date().isoformat()
            (revenue if txn['type'] == 'income' else expenses)[key] += txn['amount']
    rev_values = [revenue[k] for k in bucket_keys]
    exp_values = [expenses[k] for k in bucket_keys]
    return labels, rev_values, exp_values, [r - e for r, e in zip(rev_values, exp_values)]


def get_gross_margin(start: datetime, end: datetime) -> int:
    with db_session() as conn:
        row = conn.execute(
            '''SELECT COALESCE(SUM(amount - quantity * unit_cost_snapshot),0) AS margin
               FROM transactions WHERE type='income' AND product_id IS NOT NULL
               AND transacted_at >= ? AND transacted_at <= ?''',
            (start.isoformat(timespec='minutes'), end.isoformat(timespec='minutes')),
        ).fetchone()
    return int(row['margin'] or 0)


def get_expense_categories(start: datetime, end: datetime) -> list[dict[str, Any]]:
    with db_session() as conn:
        rows = conn.execute(
            '''SELECT category, SUM(amount) AS total FROM transactions
               WHERE type='expense' AND transacted_at >= ? AND transacted_at <= ?
               GROUP BY category ORDER BY total DESC''',
            (start.isoformat(timespec='minutes'), end.isoformat(timespec='minutes')),
        ).fetchall()
    return [{'name': r['category'], 'value': int(r['total'] or 0)} for r in rows]


def notify_error(message: str) -> None:
    ui.notify(message, type='negative', position='top')


def save_transaction(data: dict[str, Any], txn_id: int | None = None) -> None:
    try:
        occurred = datetime.fromisoformat(data['transacted_at']).replace(second=0, microsecond=0)
    except (ValueError, TypeError):
        raise ValueError('請輸入有效的日期與時間')
    amount = int(float(data['amount']))
    if amount <= 0:
        raise ValueError('金額必須大於 0')
    txn_type = data['type']
    category = (data.get('category') or '').strip()
    if not category:
        raise ValueError('請選擇帳務分類')
    product_id = data.get('product_id') if txn_type == 'income' else None
    product_name = ''
    quantity = 0.0
    unit_cost = 0
    description = (data.get('description') or '').strip()
    if product_id:
        with db_session() as conn:
            product = conn.execute('SELECT * FROM products WHERE id=? AND active=1', (int(product_id),)).fetchone()
        if product:
            product_name = product['name']
            quantity = max(float(data.get('quantity') or 1), 0)
            unit_cost = int(product['unit_cost'])
            if not description:
                description = product_name
    payload = (
        occurred.isoformat(timespec='minutes'), txn_type, category, description, amount,
        data.get('payment_method') or '現金', product_id, product_name, quantity, unit_cost,
        (data.get('notes') or '').strip(),
    )
    with db_session() as conn:
        if txn_id is None:
            conn.execute(
                '''INSERT INTO transactions
                   (transacted_at,type,category,description,amount,payment_method,product_id,
                    product_name_snapshot,quantity,unit_cost_snapshot,notes,created_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?)''',
                (*payload, datetime.now().isoformat(timespec='seconds')),
            )
        else:
            conn.execute(
                '''UPDATE transactions SET transacted_at=?,type=?,category=?,description=?,amount=?,
                   payment_method=?,product_id=?,product_name_snapshot=?,quantity=?,unit_cost_snapshot=?,notes=?
                   WHERE id=?''',
                (*payload, txn_id),
            )


def style_excel_sheet(ws) -> None:
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2563EB')
        cell.alignment = Alignment(horizontal='center')
    for column_cells in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in column_cells), default=10)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 32)


def create_excel_export() -> Path:
    txns = sorted(get_transactions(), key=lambda item: (item['transacted_at'], item['id']))
    wb = Workbook()
    ws = wb.active
    ws.title = '帳務明細'
    ws.append(['交易時間', '類型', '分類', '項目', '金額（元）', '付款方式', '商品', '數量', '單份成本（元）', '備註'])
    for t in txns:
        ws.append([
            t['transacted_at'], '收入' if t['type'] == 'income' else '支出', t['category'],
            t['description'], t['amount'], t['payment_method'], t['product_name_snapshot'],
            t['quantity'] if t['product_id'] else '', t['unit_cost_snapshot'] if t['product_id'] else '', t['notes'],
        ])
    style_excel_sheet(ws)

    daily: dict[str, dict[str, int]] = defaultdict(lambda: {'收入': 0, '支出': 0})
    monthly: dict[str, dict[str, int]] = defaultdict(lambda: {'收入': 0, '支出': 0})
    for t in txns:
        dt = parse_dt(t['transacted_at'])
        type_label = '收入' if t['type'] == 'income' else '支出'
        daily[dt.date().isoformat()][type_label] += int(t['amount'])
        monthly[dt.strftime('%Y-%m')][type_label] += int(t['amount'])
    ws_day = wb.create_sheet('每日彙總')
    ws_day.append(['日期', '收入（元）', '支出（元）', '帳務結餘（元）'])
    for day, totals in sorted(daily.items()):
        ws_day.append([day, totals['收入'], totals['支出'], totals['收入'] - totals['支出']])
    style_excel_sheet(ws_day)
    ws_month = wb.create_sheet('每月彙總')
    ws_month.append(['月份', '收入（元）', '支出（元）', '帳務結餘（元）'])
    for month, totals in sorted(monthly.items()):
        ws_month.append([month, totals['收入'], totals['支出'], totals['收入'] - totals['支出']])
    style_excel_sheet(ws_month)

    safe_name = ''.join(ch for ch in get_shop_name() if ch not in '\\/:*?"<>|').strip() or '店家'
    path = EXPORT_DIR / f'{safe_name}帳務_{datetime.now():%Y%m%d_%H%M%S}.xlsx'
    wb.save(path)
    return path


def create_db_backup() -> Path:
    path = BACKUP_DIR / f'shop_backup_{datetime.now():%Y%m%d_%H%M%S}.sqlite3'
    source = sqlite3.connect(DB_PATH)
    destination = sqlite3.connect(path)
    try:
        source.backup(destination)
    finally:
        destination.close()
        source.close()
    return path


def get_lan_ip() -> str:
    sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        sock.connect(('192.0.2.1', 80))
        return sock.getsockname()[0]
    except OSError:
        return '請查看本機區域網路 IP'
    finally:
        sock.close()


def add_head_styles() -> None:
    ui.add_head_html('''
    <style>
      :root { --brand: #16a36a; --brand-dark: #128052; --page-bg: #f1f5f9; }
      body { background: var(--page-bg); color: #172b4d; font-family: Inter, "Noto Sans TC", "Microsoft JhengHei", sans-serif; }
      .nicegui-content { padding: 0 !important; min-height: 100vh; }
      .app-card { background: white; border: 1px solid #e7edf4; border-radius: 18px; box-shadow: 0 3px 14px rgba(30, 55, 90, .035); }
      .nav-button { border-radius: 11px; min-height: 42px; }
      .nav-active { background: #e3f5eb !important; color: #11734b !important; font-weight: 600; }
      .metric-label { color: #64748b; font-size: .88rem; }
      .metric-value { color: #172b4d; font-size: clamp(1.35rem, 2vw, 1.8rem); font-weight: 750; letter-spacing: -.03em; }
      .muted { color: #718096; }
      .section-title { font-size: 1.05rem; font-weight: 700; color: #1e293b; }
      .q-field--outlined .q-field__control { border-radius: 10px; }
      .q-btn { border-radius: 10px; }
      .page-width { width: 100%; max-width: 1450px; margin: 0 auto; }
      @media (max-width: 600px) {
        .mobile-full { width: 100% !important; }
        .nicegui-content { padding-bottom: 18px !important; }
      }
    </style>
    ''')


def render_layout(active: str, title: str, subtitle: str, builder) -> None:
    shop_name = get_shop_name()
    shop_subtitle = get_shop_subtitle()
    ui.page_title(f'{title}｜{shop_name}管理系統')
    add_head_styles()
    with ui.left_drawer(value=False).props('show-if-above bordered').classes('bg-grey-1') as drawer:
        with ui.column().classes('w-full h-full p-4 gap-2'):
            with ui.row().classes('items-center gap-3 px-2 py-3 mb-2'):
                with ui.element('div').classes('w-11 h-11 rounded-xl flex items-center justify-center bg-[#16a36a] text-white'):
                    ui.icon('store', size='26px')
                with ui.column().classes('gap-0'):
                    ui.label(shop_name).classes('text-base font-bold text-slate-800')
                    ui.label(shop_subtitle).classes('text-xs text-slate-500')
            ui.separator().classes('mb-2')
            nav_items = [
                ('dashboard', '/dashboard', '儀表板', 'space_dashboard'),
                ('records', '/records', '日常記帳', 'wallet'),
                ('reports', '/reports', '統計與報表', 'query_stats'),
                ('products', '/products', '商品與成本', 'inventory_2'),
                ('history', '/history', '帳務明細', 'receipt_long'),
                ('settings', '/settings', '備份與設定', 'settings'),
            ]

            def navigate(route: str) -> None:
                # 不額外呼叫 drawer.hide()：切頁會直接建立新頁面，手機側欄由 value=False 預設收合。
                # 移除多餘的側欄動畫／元件更新，避免手機點選選單時感覺停頓；桌面仍由 show-if-above 控制。
                ui.navigate.to(route)

            for key, route, label, icon in nav_items:
                button = ui.button(label, icon=icon, on_click=lambda r=route: navigate(r)).props('flat no-caps align=left').classes('nav-button w-full justify-start')
                if key == active:
                    button.classes(add='nav-active')
            ui.space()
            ui.label('僅供店內區域網路使用').classes('text-xs text-slate-400 px-2 pb-1')
    with ui.column().classes('w-full min-h-screen gap-0'):
        with ui.row().classes('w-full items-center justify-between bg-white border-b border-slate-200 px-4 md:px-7 py-3 gap-3'):
            with ui.row().classes('items-center gap-3'):
                ui.button(icon='menu', on_click=drawer.toggle).props('flat round dense').classes('lg:hidden text-slate-600')
                with ui.column().classes('gap-0'):
                    ui.label(title).classes('text-lg md:text-xl font-bold text-slate-800')
                    ui.label(subtitle).classes('text-xs md:text-sm text-slate-500')
            with ui.row().classes('items-center gap-2'):
                ui.icon('calendar_today', color='#64748b').classes('text-sm')
                ui.label(datetime.now().strftime('%Y/%m/%d')).classes('text-sm text-slate-500')
        with ui.column().classes('page-width p-4 md:p-7 gap-5'):
            builder()


def metric_card(label: str, value: str, description: str, icon: str, accent: str, value_color: str = '#172b4d') -> None:
    with ui.card().classes('app-card p-4 md:p-5 flex-1 min-w-[145px]'):
        with ui.row().classes('w-full items-start justify-between no-wrap'):
            with ui.column().classes('gap-2'):
                ui.label(label).classes('metric-label')
                ui.label(value).classes('metric-value').style(f'color:{value_color}')
                ui.label(description).classes('text-xs text-slate-500')
            with ui.element('div').classes('w-10 h-10 rounded-xl flex items-center justify-center').style(f'background:{accent}22;color:{accent}'):
                ui.icon(icon, size='22px')


def transaction_label(t: dict[str, Any]) -> str:
    return '收入' if t['type'] == 'income' else '支出'


def render_transaction_list(txns: list[dict[str, Any]], empty_message: str = '目前沒有帳務資料') -> None:
    if not txns:
        with ui.card().classes('app-card p-8 w-full items-center'):
            ui.icon('receipt_long', size='36px').classes('text-slate-300')
            ui.label(empty_message).classes('text-slate-500')
        return
    for t in txns:
        is_income = t['type'] == 'income'
        with ui.card().classes('app-card p-3 md:p-4 w-full'):
            with ui.row().classes('w-full items-center gap-3 no-wrap'):
                with ui.element('div').classes('w-10 h-10 rounded-xl flex items-center justify-center shrink-0').style(
                    f"background:{'#dcfce7' if is_income else '#ffedd5'};color:{'#15803d' if is_income else '#c2410c'}"):
                    ui.icon('south_west' if is_income else 'north_east', size='20px')
                with ui.column().classes('gap-1 flex-1 min-w-0'):
                    with ui.row().classes('items-center gap-2 flex-wrap'):
                        ui.label(t['description'] or t['category']).classes('font-semibold text-slate-800')
                        ui.badge(t['category']).props('outline color=grey')
                    ui.label(f"{t['transacted_at'].replace('T', ' ')} · {t['payment_method']}").classes('text-xs text-slate-500')
                    if t.get('notes'):
                        ui.label(t['notes']).classes('text-xs text-slate-500 break-words')
                with ui.column().classes('items-end gap-2 shrink-0'):
                    ui.label(f"{'+' if is_income else '-'}{money(t['amount'])}").classes('font-bold text-base').style(
                        f"color:{'#15803d' if is_income else '#c2410c'}")
                    with ui.row().classes('gap-1'):
                        ui.button(icon='edit', on_click=lambda tid=t['id']: show_transaction_dialog(tid)).props('flat round dense size=sm color=grey').tooltip('編輯')
                        ui.button(icon='delete_outline', on_click=lambda tid=t['id']: confirm_delete_transaction(tid)).props('flat round dense size=sm color=grey').tooltip('刪除')


def show_transaction_dialog(txn_id: int | None = None, initial_type: str = 'income') -> None:
    txn = get_transaction(txn_id) if txn_id is not None else None
    products = get_products(active_only=True)
    product_options = {p['id']: f"{p['name']}（售價 {money(p['sale_price'])}）" for p in products}
    with ui.dialog() as dialog, ui.card().classes('w-full max-w-xl p-5 md:p-6'):
        with ui.row().classes('w-full items-center justify-between'):
            ui.label('編輯帳務' if txn else '新增帳務').classes('text-xl font-bold text-slate-800')
            ui.button(icon='close', on_click=dialog.close).props('flat round dense')
        ui.label('填寫必要欄位即可，備註與商品資料可留空。').classes('text-sm text-slate-500 mb-2')
        current_type = (txn or {}).get('type', initial_type)
        category_list = INCOME_CATEGORIES if current_type == 'income' else EXPENSE_CATEGORIES
        current_category = (txn or {}).get('category') or category_list[0]
        if current_category not in category_list:
            category_list = category_list + [current_category]
        current_product = (txn or {}).get('product_id')
        if current_product not in product_options:
            current_product = None
        type_select = ui.select({'income': '收入', 'expense': '支出'}, value=current_type, label='帳務類型').classes('w-full')
        date_input = ui.input('交易日期與時間', value=(txn or {}).get('transacted_at', datetime.now().strftime('%Y-%m-%dT%H:%M'))).props('type=datetime-local').classes('w-full')
        category_select = ui.select(category_list, value=current_category, label='分類').classes('w-full')
        description_input = ui.input('項目名稱（選填）', value=(txn or {}).get('description', ''), placeholder='例如：今日營業收入、黃豆進貨').classes('w-full')
        product_select = ui.select(product_options, value=current_product, label='商品（選填；用於預估毛利）', clearable=True).classes('w-full')
        quantity_input = ui.number('數量', value=(txn or {}).get('quantity') or 1, min=1, step=1).classes('w-full')
        amount_input = ui.number('金額（元）', value=(txn or {}).get('amount') or None, min=1, step=1).classes('w-full')
        payment_select = ui.select(PAYMENT_METHODS, value=(txn or {}).get('payment_method', '現金'), label='付款／收款方式').classes('w-full')
        notes_input = ui.textarea('備註（選填）', value=(txn or {}).get('notes', '')).props('rows=2').classes('w-full')

        def update_type(e=None) -> None:
            is_income = type_select.value == 'income'
            category_select.options = INCOME_CATEGORIES if is_income else EXPENSE_CATEGORIES
            if category_select.value not in category_select.options:
                category_select.value = category_select.options[0]
            category_select.update()
            product_select.set_visibility(is_income)
            quantity_input.set_visibility(is_income)

        def fill_from_product(e=None) -> None:
            selected = product_select.value
            if selected:
                p = next((p for p in products if p['id'] == selected), None)
                if p:
                    q = float(quantity_input.value or 1)
                    amount_input.value = int(p['sale_price'] * q)
                    if not description_input.value:
                        description_input.value = p['name']

        type_select.on_value_change(update_type)
        product_select.on_value_change(fill_from_product)
        quantity_input.on_value_change(fill_from_product)
        update_type()

        with ui.row().classes('w-full justify-end gap-2 pt-2'):
            ui.button('取消', on_click=dialog.close).props('outline color=grey no-caps')
            def submit() -> None:
                try:
                    raw_amount = amount_input.value
                    if raw_amount in (None, ''):
                        raise ValueError('請輸入金額')
                    data = {
                        'type': type_select.value,
                        'transacted_at': date_input.value,
                        'category': category_select.value,
                        'description': description_input.value,
                        'amount': raw_amount,
                        'payment_method': payment_select.value,
                        'product_id': product_select.value,
                        'quantity': quantity_input.value,
                        'notes': notes_input.value,
                    }
                    save_transaction(data, txn_id)
                    dialog.close()
                    ui.notify('帳務已儲存', type='positive', position='top')
                    ui.navigate.reload()
                except (ValueError, sqlite3.Error) as exc:
                    notify_error(str(exc) or '儲存失敗，請確認資料')
            ui.button('儲存帳務', icon='save', on_click=submit).props('unelevated no-caps color=positive')
    dialog.open()


def confirm_delete_transaction(txn_id: int) -> None:
    with ui.dialog() as dialog, ui.card().classes('w-full max-w-sm p-5'):
        ui.label('確定刪除這筆帳務？').classes('text-lg font-bold')
        ui.label('刪除後無法直接復原，請確認資料。').classes('text-sm text-slate-500')
        with ui.row().classes('w-full justify-end gap-2 pt-3'):
            ui.button('取消', on_click=dialog.close).props('outline color=grey no-caps')
            def do_delete() -> None:
                with db_session() as conn:
                    conn.execute('DELETE FROM transactions WHERE id=?', (txn_id,))
                dialog.close()
                ui.notify('帳務已刪除', type='positive')
                ui.navigate.reload()
            ui.button('確認刪除', on_click=do_delete).props('color=negative unelevated no-caps')
    dialog.open()


def show_product_dialog(product_id: int | None = None) -> None:
    products = get_products()
    product = next((p for p in products if p['id'] == product_id), None)
    with ui.dialog() as dialog, ui.card().classes('w-full max-w-md p-5 md:p-6'):
        with ui.row().classes('w-full items-center justify-between'):
            ui.label('編輯商品' if product else '新增商品').classes('text-xl font-bold')
            ui.button(icon='close', on_click=dialog.close).props('flat round dense')
        name = ui.input('商品名稱', value=(product or {}).get('name', ''), placeholder='例如：招牌甜品').classes('w-full')
        price = ui.number('售價（元）', value=(product or {}).get('sale_price') or None, min=0, step=1).classes('w-full')
        cost = ui.number('每份預估成本（元）', value=(product or {}).get('unit_cost') or 0, min=0, step=1).classes('w-full')
        active = ui.switch('啟用此商品', value=bool((product or {}).get('active', 1)))
        ui.label('成本僅用於已選取此商品的銷售紀錄之預估毛利。').classes('text-xs text-slate-500')
        with ui.row().classes('w-full justify-end gap-2 pt-2'):
            ui.button('取消', on_click=dialog.close).props('outline color=grey no-caps')
            def save_product() -> None:
                product_name = (name.value or '').strip()
                if not product_name:
                    notify_error('請輸入商品名稱')
                    return
                if price.value is None or int(float(price.value)) < 0:
                    notify_error('請輸入正確售價')
                    return
                try:
                    with db_session() as conn:
                        if product_id is None:
                            conn.execute(
                                'INSERT INTO products(name,sale_price,unit_cost,active,created_at) VALUES(?,?,?,?,?)',
                                (product_name, int(float(price.value)), int(float(cost.value or 0)), int(active.value), datetime.now().isoformat(timespec='seconds')),
                            )
                        else:
                            conn.execute(
                                'UPDATE products SET name=?,sale_price=?,unit_cost=?,active=? WHERE id=?',
                                (product_name, int(float(price.value)), int(float(cost.value or 0)), int(active.value), product_id),
                            )
                    dialog.close()
                    ui.notify('商品已儲存', type='positive')
                    ui.navigate.reload()
                except sqlite3.IntegrityError:
                    notify_error('商品名稱已存在')
            ui.button('儲存商品', icon='save', on_click=save_product).props('unelevated color=positive no-caps')
    dialog.open()


def confirm_toggle_product(product: dict[str, Any]) -> None:
    with db_session() as conn:
        conn.execute('UPDATE products SET active=? WHERE id=?', (0 if product['active'] else 1, product['id']))
    ui.notify('商品狀態已更新', type='positive')
    ui.navigate.reload()


def render_dashboard() -> None:
    now = datetime.now().replace(second=0, microsecond=0)
    today_start = now.replace(hour=0, minute=0)
    today_totals = sums_for_range(today_start, now)
    gross_margin = get_gross_margin(today_start, now)
    net = today_totals['income'] - today_totals['expense']
    with ui.row().classes('w-full items-center justify-between gap-3 flex-wrap'):
        with ui.column().classes('gap-1'):
            ui.label('今天的營運狀況').classes('text-xl font-bold text-slate-800')
            ui.label('快速掌握今日收入、支出與帳務結餘。').classes('text-sm text-slate-500')
        ui.button('＋ 新增帳務', on_click=lambda: show_transaction_dialog()).props('unelevated color=positive no-caps').classes('px-4')
    with ui.row().classes('w-full gap-3 flex-wrap'):
        metric_card('今日收入', money(today_totals['income']), '已記錄的收入', 'south_west', '#16a36a', '#15803d')
        metric_card('今日支出', money(today_totals['expense']), '已記錄的支出', 'north_east', '#f59e0b', '#c2410c')
        metric_card('預估毛利', money(gross_margin), '僅計入有選商品的銷售', 'trending_up', '#6366f1', '#4f46e5')
        metric_card('帳務結餘', money(net), '今日收入－今日支出', 'account_balance_wallet', '#0ea5a4', '#0f766e')
    with ui.row().classes('w-full items-stretch gap-4 flex-wrap'):
        with ui.card().classes('app-card p-4 md:p-5 flex-[2] min-w-[280px]'):
            with ui.row().classes('w-full items-center justify-between mb-2'):
                ui.label('最近 7 天營收與支出').classes('section-title')
                ui.button('完整報表', icon='arrow_forward', on_click=lambda: ui.navigate.to('/reports')).props('flat no-caps color=positive')
            start = (now - timedelta(days=6)).replace(hour=0, minute=0)
            labels, rev, exp, net_series = chart_series('7d', start, now)
            ui.echart({
                'color': ['#16a36a', '#f59e0b', '#64748b'],
                'tooltip': {'trigger': 'axis'},
                'legend': {'data': ['營收', '支出', '帳務結餘'], 'bottom': 0, 'textStyle': {'color': '#64748b'}},
                'grid': {'left': 10, 'right': 12, 'top': 20, 'bottom': 48, 'containLabel': True},
                'xAxis': {'type': 'category', 'data': labels, 'axisLine': {'lineStyle': {'color': '#dbe3ec'}}, 'axisLabel': {'color': '#718096'}},
                'yAxis': {'type': 'value', 'axisLabel': {'color': '#718096', 'formatter': '${value}'}, 'splitLine': {'lineStyle': {'color': '#edf1f6'}}},
                'series': [
                    {'name': '營收', 'type': 'line', 'smooth': True, 'symbolSize': 6, 'data': rev},
                    {'name': '支出', 'type': 'line', 'smooth': True, 'symbolSize': 6, 'data': exp},
                    {'name': '帳務結餘', 'type': 'line', 'smooth': True, 'symbolSize': 6, 'data': net_series},
                ],
            }).classes('w-full h-72')
        with ui.card().classes('app-card p-4 md:p-5 flex-1 min-w-[260px]'):
            ui.label('常用功能').classes('section-title mb-2')
            quick = [
                ('新增收入／支出', 'add_circle', lambda: show_transaction_dialog()),
                ('查看帳務明細', 'receipt_long', lambda: ui.navigate.to('/history')),
                ('分析利潤趨勢', 'query_stats', lambda: ui.navigate.to('/reports')),
                ('匯出 Excel', 'file_download', export_excel),
            ]
            for label, icon, callback in quick:
                ui.button(label, icon=icon, on_click=callback).props('outline no-caps align=left').classes('w-full justify-start mb-1')
    with ui.card().classes('app-card p-4 md:p-5 w-full'):
        with ui.row().classes('w-full items-center justify-between mb-2'):
            ui.label('最近帳務').classes('section-title')
            ui.button('查看全部', on_click=lambda: ui.navigate.to('/history')).props('flat no-caps color=positive')
        render_transaction_list(get_transactions(limit=5), '目前尚無帳務，點選「新增帳務」開始記錄。')


def render_records() -> None:
    with ui.row().classes('w-full items-center justify-between gap-3 flex-wrap'):
        with ui.column().classes('gap-1'):
            ui.label('日常記帳').classes('text-xl font-bold text-slate-800')
            ui.label('可快速新增收入或支出，也能修正已輸入的紀錄。').classes('text-sm text-slate-500')
        with ui.row().classes('gap-2 flex-wrap'):
            ui.button('新增收入', icon='add', on_click=lambda: show_transaction_dialog(initial_type='income')).props('unelevated color=positive no-caps')
            ui.button('新增支出', icon='remove', on_click=lambda: show_transaction_dialog(initial_type='expense')).props('outline color=deep-orange no-caps')
    now = datetime.now().replace(second=0, microsecond=0)
    today_start = now.replace(hour=0, minute=0)
    totals = sums_for_range(today_start, now)
    with ui.row().classes('w-full gap-3 flex-wrap'):
        metric_card('今日收入', money(totals['income']), '今天已記錄', 'south_west', '#16a36a', '#15803d')
        metric_card('今日支出', money(totals['expense']), '今天已記錄', 'north_east', '#f59e0b', '#c2410c')
        metric_card('今日結餘', money(totals['income'] - totals['expense']), '收入減支出', 'account_balance_wallet', '#0ea5a4', '#0f766e')
    with ui.card().classes('app-card p-4 md:p-5 w-full'):
        ui.label('最近 20 筆帳務').classes('section-title mb-3')
        render_transaction_list(get_transactions(limit=20))


def render_reports() -> None:
    ui.label('依時段查看營收、支出及帳務結餘。').classes('text-sm text-slate-500')
    with ui.card().classes('app-card p-4 md:p-5 w-full'):
        with ui.row().classes('w-full items-end gap-3 flex-wrap'):
            period_select = ui.select({
                '24h': '最近 24 小時', '7d': '最近 7 天', '30d': '最近 30 天',
                'month': '本月', 'year': '今年', 'custom': '自訂日期',
            }, value='7d', label='統計期間').classes('w-full sm:w-52')
            start_input = ui.input('開始日期', value=(date.today() - timedelta(days=6)).isoformat()).props('type=date').classes('w-full sm:w-48')
            end_input = ui.input('結束日期', value=date.today().isoformat()).props('type=date').classes('w-full sm:w-48')
            apply_button = ui.button('套用篩選', icon='filter_alt').props('unelevated color=positive no-caps')
        summary = ui.row().classes('w-full gap-3 flex-wrap mt-4')
        chart_box = ui.column().classes('w-full gap-4 mt-2')
        expense_box = ui.column().classes('w-full gap-4')

        def refresh() -> None:
            start, end, label = period_bounds(period_select.value, start_input.value, end_input.value)
            totals = sums_for_range(start, end)
            net = totals['income'] - totals['expense']
            summary.clear()
            with summary:
                metric_card('期間收入', money(totals['income']), label, 'south_west', '#16a36a', '#15803d')
                metric_card('期間支出', money(totals['expense']), label, 'north_east', '#f59e0b', '#c2410c')
                metric_card('帳務結餘', money(net), '收入－支出（簡易現金口徑）', 'account_balance_wallet', '#0ea5a4', '#0f766e')
                metric_card('預估毛利', money(get_gross_margin(start, end)), '限有商品成本資料的銷售', 'trending_up', '#6366f1', '#4f46e5')
            labels, rev, exp, net_values = chart_series(period_select.value, start, end)
            chart_box.clear()
            with chart_box:
                with ui.card().classes('app-card p-4 md:p-5 w-full'):
                    ui.label(f'{label}｜營收、支出與帳務結餘').classes('section-title')
                    ui.echart({
                        'color': ['#16a36a', '#f59e0b', '#64748b'],
                        'tooltip': {'trigger': 'axis'},
                        'legend': {'data': ['營收', '支出', '帳務結餘'], 'bottom': 0, 'textStyle': {'color': '#64748b'}},
                        'grid': {'left': 10, 'right': 15, 'top': 25, 'bottom': 48, 'containLabel': True},
                        'xAxis': {'type': 'category', 'data': labels, 'axisLabel': {'color': '#718096', 'hideOverlap': True}, 'axisLine': {'lineStyle': {'color': '#dbe3ec'}}},
                        'yAxis': {'type': 'value', 'axisLabel': {'color': '#718096', 'formatter': '${value}'}, 'splitLine': {'lineStyle': {'color': '#edf1f6'}}},
                        'series': [
                            {'name': '營收', 'type': 'line', 'smooth': True, 'symbolSize': 5, 'data': rev},
                            {'name': '支出', 'type': 'line', 'smooth': True, 'symbolSize': 5, 'data': exp},
                            {'name': '帳務結餘', 'type': 'line', 'smooth': True, 'symbolSize': 5, 'data': net_values},
                        ],
                    }).classes('w-full h-80')
            expense_box.clear()
            categories = get_expense_categories(start, end)
            with expense_box:
                if categories:
                    with ui.card().classes('app-card p-4 md:p-5 w-full'):
                        ui.label('支出分類').classes('section-title')
                        ui.echart({
                            'tooltip': {'trigger': 'item', 'formatter': '{b}: ${c} ({d}%)'},
                            'legend': {'bottom': 0, 'type': 'scroll'},
                            'series': [{'name': '支出分類', 'type': 'pie', 'radius': ['40%', '68%'], 'center': ['50%', '45%'],
                                        'avoidLabelOverlap': True,
                                        'itemStyle': {'borderRadius': 6, 'borderColor': '#fff', 'borderWidth': 2},
                                        'label': {'formatter': '{b}\n${c}'}, 'data': categories}],
                        }).classes('w-full h-72')
                else:
                    ui.label('這個期間尚無支出資料。').classes('text-sm text-slate-500')

        def period_changed(e=None) -> None:
            custom = period_select.value == 'custom'
            start_input.set_visibility(custom)
            end_input.set_visibility(custom)
            if not custom:
                start, end, _ = period_bounds(period_select.value)
                start_input.value = start.date().isoformat()
                end_input.value = end.date().isoformat()
            refresh()

        period_select.on_value_change(period_changed)
        apply_button.on_click(refresh)
        start_input.set_visibility(False)
        end_input.set_visibility(False)
        refresh()
        with ui.row().classes('w-full justify-end pt-1'):
            ui.button('匯出全部帳務 Excel', icon='file_download', on_click=export_excel).props('outline color=positive no-caps')


def render_products() -> None:
    with ui.row().classes('w-full items-center justify-between gap-3 flex-wrap'):
        with ui.column().classes('gap-1'):
            ui.label('商品與成本').classes('text-xl font-bold text-slate-800')
            ui.label('設定商品售價與每份預估成本，供銷售紀錄估算毛利。').classes('text-sm text-slate-500')
        ui.button('新增商品', icon='add', on_click=lambda: show_product_dialog()).props('unelevated color=positive no-caps')
    products = get_products()
    if not products:
        with ui.card().classes('app-card w-full p-8 items-center'):
            ui.icon('inventory_2', size='40px').classes('text-slate-300')
            ui.label('尚未建立商品').classes('text-lg font-semibold text-slate-700')
            ui.label('新增常賣的甜品或配料商品後，記帳時即可選取。').classes('text-sm text-slate-500 text-center')
            ui.button('建立第一個商品', on_click=lambda: show_product_dialog()).props('unelevated color=positive no-caps')
        return
    with ui.grid(columns=1).classes('w-full gap-3'):
        for p in products:
            with ui.card().classes('app-card p-4 w-full'):
                with ui.row().classes('w-full items-center gap-3 no-wrap'):
                    with ui.element('div').classes('w-11 h-11 rounded-xl bg-emerald-50 text-emerald-700 flex items-center justify-center shrink-0'):
                        ui.icon('icecream', size='24px')
                    with ui.column().classes('gap-1 flex-1 min-w-0'):
                        with ui.row().classes('items-center gap-2 flex-wrap'):
                            ui.label(p['name']).classes('font-semibold text-slate-800')
                            ui.badge('啟用' if p['active'] else '已停用', color='positive' if p['active'] else 'grey')
                        ui.label(f"售價 {money(p['sale_price'])}　·　預估成本 {money(p['unit_cost'])}　·　單份毛利 {money(p['sale_price'] - p['unit_cost'])}").classes('text-sm text-slate-500')
                    with ui.row().classes('items-center gap-1'):
                        ui.button(icon='edit', on_click=lambda pid=p['id']: show_product_dialog(pid)).props('flat round color=grey')
                        ui.button(icon='toggle_on' if p['active'] else 'toggle_off', on_click=lambda item=p: confirm_toggle_product(item)).props('flat round color=grey').tooltip('切換啟用狀態')


def render_history() -> None:
    with ui.row().classes('w-full items-center justify-between gap-3 flex-wrap'):
        with ui.column().classes('gap-1'):
            ui.label('帳務明細').classes('text-xl font-bold text-slate-800')
            ui.label('依日期、類型與關鍵字查詢，並可編輯或刪除紀錄。').classes('text-sm text-slate-500')
        with ui.row().classes('gap-2 flex-wrap'):
            ui.button('匯出 Excel', icon='file_download', on_click=export_excel).props('outline color=positive no-caps')
            ui.button('新增帳務', icon='add', on_click=lambda: show_transaction_dialog()).props('unelevated color=positive no-caps')
    with ui.card().classes('app-card p-4 md:p-5 w-full'):
        with ui.row().classes('w-full items-end gap-3 flex-wrap'):
            start_input = ui.input('開始日期').props('type=date').classes('w-full sm:w-48')
            end_input = ui.input('結束日期').props('type=date').classes('w-full sm:w-48')
            type_select = ui.select({'all': '全部類型', 'income': '收入', 'expense': '支出'}, value='all', label='類型').classes('w-full sm:w-40')
            keyword = ui.input('關鍵字', placeholder='分類、項目或備註').classes('w-full sm:flex-1')
            query_button = ui.button('查詢', icon='search').props('unelevated color=positive no-caps')
        results = ui.column().classes('w-full gap-3 mt-3')
        def search() -> None:
            start = None
            end = None
            try:
                if start_input.value:
                    start = datetime.combine(date.fromisoformat(start_input.value), datetime.min.time())
                if end_input.value:
                    end = datetime.combine(date.fromisoformat(end_input.value), datetime.max.time()).replace(second=0, microsecond=0)
            except ValueError:
                notify_error('日期格式不正確')
                return
            if start and end and end < start:
                notify_error('結束日期不能早於開始日期')
                return
            txns = get_transactions(start, end)
            if type_select.value != 'all':
                txns = [t for t in txns if t['type'] == type_select.value]
            query = (keyword.value or '').strip().lower()
            if query:
                txns = [t for t in txns if query in ' '.join(str(t.get(k) or '') for k in ('category', 'description', 'notes', 'payment_method', 'product_name_snapshot')).lower()]
            results.clear()
            with results:
                ui.label(f'查詢結果：{len(txns)} 筆').classes('text-sm text-slate-500')
                render_transaction_list(txns, '沒有符合條件的帳務資料')
        query_button.on_click(search)
        start_input.value = (date.today() - timedelta(days=30)).isoformat()
        end_input.value = date.today().isoformat()
        search()


def export_excel() -> None:
    try:
        path = create_excel_export()
        ui.download(path.read_bytes(), filename=path.name)
        ui.notify('Excel 報表已準備下載', type='positive', position='top')
    except Exception as exc:
        notify_error(f'匯出失敗：{exc}')


def backup_database() -> None:
    try:
        path = create_db_backup()
        ui.download(path.read_bytes(), filename=path.name)
        ui.notify('資料庫備份已準備下載', type='positive', position='top')
    except Exception as exc:
        notify_error(f'備份失敗：{exc}')


def render_settings() -> None:
    lan_ip = get_lan_ip()
    with ui.card().classes('app-card p-5 md:p-6 w-full'):
        ui.label('店家名稱').classes('section-title')
        ui.label('設定顯示於側欄、頁面標題與匯出檔名的店名，方便套用到不同店家。').classes('text-sm text-slate-500')
        name_input = ui.input('店家名稱', value=get_shop_name(), placeholder='例如：豆花冰店').classes('w-full mt-3')
        subtitle_input = ui.input('副標題（選填）', value=get_shop_subtitle(), placeholder='例如：營運管理系統').classes('w-full')

        def save_shop_name() -> None:
            new_name = (name_input.value or '').strip()
            if not new_name:
                notify_error('請輸入店家名稱')
                return
            new_subtitle = (subtitle_input.value or '').strip() or DEFAULT_SHOP_SUBTITLE
            set_setting('shop_name', new_name)
            set_setting('shop_subtitle', new_subtitle)
            ui.notify('店家名稱已更新', type='positive', position='top')
            ui.navigate.reload()

        with ui.row().classes('w-full justify-end pt-2'):
            ui.button('儲存店家名稱', icon='save', on_click=save_shop_name).props('unelevated color=positive no-caps')
    with ui.card().classes('app-card p-5 md:p-6 w-full'):
        ui.label('店內存取方式').classes('section-title')
        ui.label('請讓手機連上與店內電腦相同的 Wi-Fi，再於手機瀏覽器輸入以下網址。').classes('text-sm text-slate-500')
        with ui.row().classes('items-center gap-2 flex-wrap mt-3'):
            with ui.element('div').classes('bg-slate-100 rounded-xl px-4 py-3'):
                ui.label(f'http://{lan_ip}:{PORT}').classes('font-mono text-base md:text-lg font-semibold text-slate-800')
        ui.label('若網址無法開啟，請確認電腦與手機在同一個區域網路，且 Windows 防火牆允許 TCP 連接埠。').classes('text-xs text-slate-500 mt-2')
    with ui.card().classes('app-card p-5 md:p-6 w-full'):
        ui.label('資料備份').classes('section-title')
        ui.label('備份會產生一份 SQLite 資料庫檔案，包含目前帳務與商品資料。建議定期下載並存放在另一個安全位置。').classes('text-sm text-slate-500')
        with ui.row().classes('gap-2 flex-wrap mt-3'):
            ui.button('下載資料庫備份', icon='database', on_click=backup_database).props('unelevated color=positive no-caps')
            ui.button('匯出 Excel 帳務', icon='file_download', on_click=export_excel).props('outline color=positive no-caps')
    with ui.card().classes('app-card p-5 md:p-6 w-full'):
        ui.label('本機部署注意事項').classes('section-title')
        for text in [
            '店內電腦必須保持開機，且管理程式持續執行，手機才能使用。',
            '不要在路由器設定連接埠轉發；Windows 防火牆僅允許私人網路存取 TCP 8080。',
            '同一區域網路中的裝置可能連到管理頁面，請使用可信任的店內 Wi-Fi。',
            f'資料庫位置：{DB_PATH}',
        ]:
            with ui.row().classes('items-start gap-2'):
                ui.icon('check_circle', color='#16a36a').classes('mt-1')
                ui.label(text).classes('text-sm text-slate-600 break-all')


@ui.page('/')
def index_page() -> None:
    ui.navigate.to('/dashboard')


@ui.page('/dashboard')
def dashboard_page() -> None:
    render_layout('dashboard', '營運儀表板', '店內營運概況一覽', render_dashboard)


@ui.page('/records')
def records_page() -> None:
    render_layout('records', '日常記帳', '快速新增與管理收入支出', render_records)


@ui.page('/reports')
def reports_page() -> None:
    render_layout('reports', '統計與報表', '依時段查看營運趨勢與支出結構', render_reports)


@ui.page('/products')
def products_page() -> None:
    render_layout('products', '商品與成本', '管理商品售價與預估成本', render_products)


@ui.page('/history')
def history_page() -> None:
    render_layout('history', '帳務明細', '查詢、修改及匯出歷史帳務', render_history)


@ui.page('/settings')
def settings_page() -> None:
    render_layout('settings', '備份與設定', '本地網路使用方式與資料保護', render_settings)


def _silence_connection_reset() -> None:
    """忽略瀏覽器中斷連線時 Windows asyncio 丟出的無害 ConnectionResetError 噪音。"""
    try:
        loop = asyncio.get_running_loop()
    except RuntimeError:
        loop = asyncio.get_event_loop()
    previous = loop.get_exception_handler()

    def handler(target_loop: asyncio.AbstractEventLoop, context: dict[str, Any]) -> None:
        exc = context.get('exception')
        if isinstance(exc, ConnectionResetError):
            return  # 瀏覽器端關閉連線，屬正常情況
        if previous is not None:
            previous(target_loop, context)
        else:
            target_loop.default_exception_handler(context)

    loop.set_exception_handler(handler)


web_app.on_startup(_silence_connection_reset)


if __name__ in {'__main__', '__mp_main__'}:
    ui.run(
        title=f'{get_shop_name()}管理系統',
        host='0.0.0.0',
        port=PORT,
        reload=False,
        show=False,
        favicon='🍧',
    )
